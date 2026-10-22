"""Golden de exportación: la propuesta de Alborada sobre el kit se renderiza a Word (y PDF si hay
LibreOffice). Guarda el ejemplo en `docs/ejemplos/propuesta_alborada_ejemplo.docx`."""

import zipfile
from datetime import datetime
from io import BytesIO
from pathlib import Path

import pytest
from docx import Document

from posiq.adapters.export import (
    AgencyBrandKit,
    DatosDiferencial,
    DocxExporter,
    PdfExporter,
    SofficeNotFoundError,
)
from posiq.adapters.export.diferencial import (
    AUDITORIA_FIXTURE,
    ESCENARIOS_INTRO,
    NO_CALCULADO,
    PILARES_SEGUIMIENTO,
)
from posiq.adapters.export.docx import ROTULO_SINTETICO, fmt_cifra, fmt_score
from posiq.adapters.export.pdf import find_soffice
from posiq.adapters.export.template import SECCIONES, TEMPLATE_PATH, build_template
from posiq.adapters.web import FixtureWebFetcher
from posiq.application.nodes.diferencial import calcular, escenarios_del_plan
from posiq.domain.brief_coverage import evaluar_cobertura, puntos_del_brief
from posiq.domain.findings import detect_findings
from posiq.domain.metrics import compute_metrics
from posiq.domain.schemas import (
    BriefSummary,
    Competencia,
    Contexto,
    Diagnostico,
    Estrategia,
    Hito,
    HojaRuta,
    Medicion,
    Objective,
    Proposal,
    Riesgo,
    UnidadCifra,
)
from posiq.domain.simulation import simulate_plan
from posiq.domain.web_audit import analizar_web
from posiq.domain.web_audit_schemas import RECURSOS, AuditoriaWeb
from posiq.ports import Exporter
from tests.golden.conftest import STUDY_ID, Kit, example_plan

EJEMPLO = (
    Path(__file__).resolve().parents[2] / "docs" / "ejemplos" / "propuesta_alborada_ejemplo.docx"
)


@pytest.fixture(scope="module")
def proposal(kit: Kit, alborada_brief: BriefSummary) -> Proposal:
    metrics = compute_metrics(
        alborada_brief,
        kit.market,
        kit.prompts,
        kit.runs,
        kit.evaluations(),
        kit.sources,
        study_id=STUDY_ID,
    )
    findings = detect_findings(metrics, alborada_brief)
    # El plan canónico de Alborada (9 acciones en 4 fases): el de la simulación 25,9 → 30,8.
    plan = example_plan(findings, alborada_brief)
    responsabilidades: dict[str, list[str]] = {}
    for a in plan.acciones:
        responsabilidades.setdefault(a.responsable, []).append(a.id)
    return Proposal(
        study_id=STUDY_ID,
        resumen_ejecutivo=(
            "Clínica Alborada Visión cuenta con un sólido reconocimiento asistido en los "
            "asistentes de IA analizados, alcanzando una presencia del 100 % cuando los usuarios "
            "consultan directamente por la marca.\n"
            "Sin embargo, ante consultas abiertas y espontáneas sobre cirugía refractiva y control "
            "de miopía infantil, la recomendación desciende significativamente, situándose por "
            "detrás del líder del sector (Grupo Hospitalario Valdecruz).\n"
            "Se propone un plan estratégico por fases orientado a reforzar la presencia en fuentes "
            "sanitarias de referencia, generar contenido especializado para momentos de decisión y "
            "recuperar visibilidad en Barcelona, dentro de las restricciones del brief y los "
            "recursos del equipo; cada contenido requiere revisión médica y jurídica antes de "
            "publicarse."
        ),
        contexto=Contexto(
            marca=alborada_brief.marca,
            sector=alborada_brief.sector,
            motores=["ChatGPT", "Gemini", "Perplexity", "Claude"],
            fecha=datetime(2026, 10, 21),
            origen_datos="kit",
            datos_sinteticos=True,
            limitaciones=[
                "Datos sintéticos del kit del reto: reproducen con fidelidad la estructura de "
                "un estudio GEO, pero no reflejan cuotas reales de mercado.",
                "Naturaleza probabilística de los modelos de lenguaje: las respuestas varían en "
                "el tiempo y requieren seguimiento periódico.",
                "Publicidad sanitaria regulada (RD 1907/1996): el plan no promete resultados ni "
                "hace comparativas denigratorias, y cada contenido necesita revisión médica y "
                "jurídica de la clínica antes de publicarse. Las mejoras de visibilidad son "
                "hipótesis técnicas y no garantizan resultados clínicos ni comerciales.",
            ],
        ),
        diagnostico=Diagnostico(
            texto="Diagnóstico detallado sobre la visibilidad espontánea y asistida por motor, "
            "fase del embudo y ciudad.",
            metrics=metrics,
        ),
        competencia=Competencia(
            texto="Análisis comparativo de posicionamiento y brecha de fuentes frente al líder "
            "del mercado y competidores.",
            ranking=metrics.ranking,
        ),
        hallazgos=findings,
        estrategia=Estrategia(
            texto="Estrategia orientada a consolidar la prescripción espontánea en Madrid y "
            "Barcelona, impulsando la autoridad en fuentes médicas clave.",
            objetivos=[
                Objective(
                    kpi="Score de Citación",
                    valor_actual=metrics.score_global.score or 0,
                    valor_objetivo=35,
                    plazo_semanas=26,
                )
            ],
            prioridades=["Autoridad en terceros", "Contenido por ciudad"],
        ),
        plan=plan,
        medicion=Medicion(
            kpis=["Visibilidad espontánea", "Score de Citación"],
            metodologia="Re-ejecución del estudio con el mismo banco de prompts",
            cadencia="Trimestral",
        ),
        riesgos=[
            Riesgo(
                riesgo="Variabilidad de fuentes y modelos de IA",
                mitigacion="Protocolo de re-medición periódica para monitorizar cambios "
                "algorítmicos",
            ),
            Riesgo(
                riesgo="Marco regulatorio en publicidad sanitaria",
                mitigacion="Revisión médica previa de todo contenido divulgativo sin promesas de "
                "resultados",
            ),
        ],
        hoja_ruta=HojaRuta(
            hitos=[
                Hito(semana=0, descripcion="Medición inicial (línea base)"),
                Hito(semana=12, descripcion="Primera re-medición trimestral comparable"),
            ],
            responsabilidades=responsabilidades,
            presupuesto_eur=22000,
            seguimiento="Comité mensual de seguimiento y presentación de informe evolutivo al "
            "comité de dirección",
        ),
        marcas={b.id: b.nombre for b in kit.brands},
    )


@pytest.fixture(scope="module")
def proposal_con_encargo(proposal: Proposal, alborada_brief: BriefSummary) -> Proposal:
    """La misma propuesta con el encargo del brief: todo enlazado salvo los pilotos."""
    acciones = [a.id for a in proposal.plan.acciones]
    enlaces = {p.id: acciones for p in puntos_del_brief(alborada_brief) if p.id != "publico-4"}
    encargo = evaluar_cobertura(alborada_brief, proposal.plan, proposal.hallazgos, enlaces)
    return proposal.model_copy(update={"encargo": encargo})


@pytest.fixture(scope="module")
def diferencial(kit: Kit, alborada_brief: BriefSummary, proposal: Proposal) -> DatosDiferencial:
    """Lo que el grafo deja en el estado para Alborada: nodo `diferencial`, escenarios y
    simulación del plan (mismas funciones que los nodos)."""
    metrics = proposal.diagnostico.metrics
    estado = {
        "study_id": STUDY_ID,
        "brief": alborada_brief,
        "market": kit.market,
        "prompts": kit.prompts,
        "runs": kit.runs,
        "evaluations": kit.evaluations(),
        "metrics": metrics,
        "findings": proposal.hallazgos,
    }
    calculado = calcular(estado, kit.sources)  # type: ignore[arg-type]
    assert calculado["errores"] == []
    escenarios, errores = escenarios_del_plan(proposal.plan, proposal.hallazgos, metrics)
    assert errores == []
    simulacion = simulate_plan(proposal.plan, metrics, proposal.hallazgos)
    return DatosDiferencial.del_estado(
        {**calculado, "escenarios": escenarios, "simulacion": simulacion}
    )


@pytest.fixture(scope="module")
def docx_bytes(proposal_con_encargo: Proposal, diferencial: DatosDiferencial) -> bytes:
    return DocxExporter().render(proposal_con_encargo, formato="docx", diferencial=diferencial)


def _texto(contenido: bytes) -> str:
    doc = Document(BytesIO(contenido))
    partes = [p.text for p in doc.paragraphs]
    partes += [c.text for t in doc.tables for row in t.rows for c in row.cells]
    partes += [p.text for s in doc.sections for p in s.header.paragraphs + s.footer.paragraphs]
    return "\n".join(partes)


def _partes(contenido: bytes) -> dict[str, bytes]:
    with zipfile.ZipFile(BytesIO(contenido)) as z:
        return {n: z.read(n) for n in z.namelist()}


def test_cumple_el_puerto() -> None:
    exporters: list[Exporter] = [DocxExporter(), PdfExporter()]
    assert all(callable(e.render) for e in exporters)


def test_docx_con_las_secciones_y_cifras(proposal: Proposal, docx_bytes: bytes) -> None:
    texto = _texto(docx_bytes)
    posiciones = [texto.index(s) for s in SECCIONES]
    assert posiciones == sorted(posiciones)
    assert "{{" not in texto and "{%" not in texto
    score = proposal.diagnostico.metrics.score_global.score
    assert fmt_score(score) in texto
    for hallazgo in proposal.hallazgos:
        assert hallazgo.titulo in texto and hallazgo.origen_cifra in texto
        assert fmt_cifra(hallazgo.cifra, hallazgo.unidad) in texto
    assert {"fraccion", "puntos", "recuento"} <= {h.unidad for h in proposal.hallazgos}
    assert "Láser Nítida" in texto and "Instituto Oftálmico Tramuntana" in texto
    for limitacion in proposal.contexto.limitaciones:
        assert limitacion in texto
    assert "posiq." in texto
    # Solo se reescribe si cambia el contenido: el zip lleva marcas de tiempo distintas cada vez.
    if not EJEMPLO.is_file() or _partes(EJEMPLO.read_bytes()) != _partes(docx_bytes):
        EJEMPLO.parent.mkdir(parents=True, exist_ok=True)
        EJEMPLO.write_bytes(docx_bytes)


def test_docx_con_lo_que_nos_habeis_pedido(
    proposal_con_encargo: Proposal, docx_bytes: bytes
) -> None:
    texto = _texto(docx_bytes)
    encargo = proposal_con_encargo.encargo
    assert encargo is not None
    inicio = texto.index(SECCIONES[1])
    assert inicio < texto.index(SECCIONES[2]) < texto.index(SECCIONES[3])
    seccion = texto[inicio : texto.index(SECCIONES[2])]
    # Los puntos literales, en viñetas dentro de la sección.
    for punto in encargo.cobertura:
        assert punto.texto in seccion
    assert "18.000–27.000 € al año" in seccion
    # La tabla de cobertura (las tablas se leen aparte de los párrafos).
    filas = [
        [c.text for c in fila.cells]
        for t in Document(BytesIO(docx_bytes)).tables
        for fila in t.rows
        if t.rows[0].cells[0].text == "Punto"
    ]
    assert filas[0] == ["Punto", "Lo que pedisteis", "Dónde se trata", "Estado"]
    por_texto = {f[1]: f for f in filas[1:]}
    assert len(por_texto) == len(encargo.cobertura)
    acciones = "Acciones " + ", ".join(a.id for a in proposal_con_encargo.plan.acciones)
    assert por_texto["Madrid"] == ["Ciudad", "Madrid", acciones, "Cubierto"]
    restriccion = por_texto["Sin comparativas denigratorias"]
    assert restriccion[2] == "La declaran todas las acciones del plan."
    presupuesto = por_texto["18.000–27.000 € al año"]
    assert presupuesto[2] == "Plan de 22.000 €, dentro del máximo de 27.000 €."
    pilotos = por_texto["Profesionales con requisitos específicos (pilotos)"]
    assert pilotos[2].startswith("pendiente: Ninguna acción ni hallazgo")
    assert [f[3] for f in filas[1:]].count("Pendiente") == 1


def test_docx_sin_encargo(proposal: Proposal) -> None:
    texto = _texto(DocxExporter().render(proposal, formato="docx"))
    assert SECCIONES[1] in texto and "no registra el encargo" in texto


@pytest.mark.parametrize(
    ("cifra", "unidad", "esperado"),
    [
        (0.15, "fraccion", "15,0 %"),
        (40.53, "puntos", "40,53 puntos"),
        (1.0, "puntos", "1 punto"),
        (37.0, "recuento", "37"),
        (22000.0, "eur", "22.000 €"),
    ],
)
def test_cifra_segun_unidad(cifra: float, unidad: UnidadCifra, esperado: str) -> None:
    assert fmt_cifra(cifra, unidad) == esperado


def test_sin_marcas_cae_a_los_ids(proposal: Proposal) -> None:
    sin_marcas = proposal.model_copy(update={"marcas": {}})
    texto = _texto(DocxExporter().render(sin_marcas, formato="docx"))
    assert "Laser Nitida" in texto and "Láser Nítida" not in texto
    assert proposal.contexto.marca in texto


def test_cabecera_con_marca_de_agencia(proposal: Proposal) -> None:
    kit = AgencyBrandKit(
        nombre="Agencia Norte", color_primario="#AA3300", pie_de_pagina="Confidencial"
    )
    texto = _texto(DocxExporter().render(proposal, formato="docx", marca_agencia=kit))
    assert "Agencia Norte" in texto and "Confidencial" in texto


def test_docx_exporter_rechaza_pdf(proposal: Proposal) -> None:
    with pytest.raises(ValueError, match="PdfExporter"):
        DocxExporter().render(proposal, formato="pdf")


def test_plantilla_reproducible(tmp_path: Path) -> None:
    nueva = build_template(tmp_path / "propuesta.docx")
    partes = ("word/document.xml", "word/header1.xml", "word/footer1.xml")
    with zipfile.ZipFile(TEMPLATE_PATH) as a, zipfile.ZipFile(nueva) as b:
        for parte in partes:
            assert a.read(parte) == b.read(parte), f"plantilla desactualizada: {parte}"


def test_pdf_sin_libreoffice_da_error_claro(
    proposal: Proposal, monkeypatch: pytest.MonkeyPatch
) -> None:
    monkeypatch.setattr("posiq.adapters.export.pdf.find_soffice", lambda: None)
    with pytest.raises(SofficeNotFoundError, match="LibreOffice"):
        PdfExporter().render(proposal, formato="pdf")


@pytest.mark.skipif(find_soffice() is None, reason="LibreOffice (soffice) no está instalado")
def test_pdf_con_libreoffice(proposal: Proposal) -> None:
    pdf = PdfExporter().render(proposal, formato="pdf")
    assert pdf.startswith(b"%PDF")


# --- Diferencial: rigor, competidores, fuentes, exactitud, escenarios y seguimiento --------------


def _seccion(texto: str, i: int) -> str:
    inicio = texto.index(SECCIONES[i])
    fin = texto.index(SECCIONES[i + 1]) if i + 1 < len(SECCIONES) else len(texto)
    return texto[inicio:fin]


def _tablas(contenido: bytes) -> list[list[list[str]]]:
    return [
        [[c.text for c in fila.cells] for fila in t.rows]
        for t in Document(BytesIO(contenido)).tables
    ]


def _tabla_con(contenido: bytes, *cabecera: str) -> list[list[str]]:
    return next(t for t in _tablas(contenido) if t[0][: len(cabecera)] == list(cabecera))


def test_caso_sintetico_en_portada_y_pie(docx_bytes: bytes) -> None:
    doc = Document(BytesIO(docx_bytes))
    portada = "\n".join(p.text for p in doc.paragraphs[:8])
    assert ROTULO_SINTETICO in portada
    pie = "\n".join(p.text for s in doc.sections for p in s.footer.paragraphs)
    assert ROTULO_SINTETICO in pie


def test_estudio_real_sin_rotulo_sintetico(proposal: Proposal) -> None:
    contexto = proposal.contexto.model_copy(
        update={"datos_sinteticos": False, "origen_datos": "api"}
    )
    real = proposal.model_copy(update={"contexto": contexto})
    assert ROTULO_SINTETICO not in _texto(DocxExporter().render(real, formato="docx"))


def test_rigor_con_ic_y_los_cuatro_motores(
    docx_bytes: bytes, diferencial: DatosDiferencial
) -> None:
    texto = _seccion(_texto(docx_bytes), 3)
    assert "Margen de error de las cifras" in texto
    assert "25,9 [21,6–30,1]" in _texto(
        docx_bytes
    )  # las fichas se leen aparte  # cifra canónica del golden de rigor
    filas = _tabla_con(docx_bytes, "Motor", "Aparición [IC 95 %]")
    assert [f[0] for f in filas[1:]] == ["ChatGPT", "Gemini", "Perplexity", "Claude"]
    assert diferencial.rigor is not None
    chatgpt = diferencial.rigor.aparicion_por_motor["ChatGPT"]
    assert filas[1][1] == "18,3 % [10,6 %–29,9 %]"
    assert filas[1][2].startswith("± ") and filas[1][3] == str(chatgpt.n)


def test_hallazgos_con_su_contraste(docx_bytes: bytes) -> None:
    contrastes = [f[1] for t in _tablas(docx_bytes) for f in t if f[0] == "Contraste estadístico"]
    assert len(contrastes) == 7  # uno por hallazgo
    texto = "\n".join(contrastes)
    assert "Se distingue del ruido de muestreo" in texto
    # F-fuente-lider no admite contraste: lo dice, no se lee como «no significativo».
    assert "Sin contraste estadístico" in texto


def test_competidores_declarados_frente_a_la_ia(docx_bytes: bytes) -> None:
    texto = _seccion(_texto(docx_bytes), 4)
    assert "Competidores declarados frente a los que pone la IA" in texto
    assert "es la 8.ª de 10 marcas" in texto  # cifra canónica: 8.ª de 10
    declarados, por_delante = [t for t in _tablas(docx_bytes) if t[0][:2] == ["Marca", "Posición"]]
    assert {f[0] for f in declarados[1:]} == {
        "Grupo Hospitalario Valdecruz",
        "Instituto Oftálmico Tramuntana",
        "Láser Nítida",
    }
    oftalmos = next(f for f in por_delante[1:] if f[0] == "Oftalmos Centro")
    assert oftalmos[5] == "Relevante"


def test_mapa_de_fuentes_y_oportunidades(docx_bytes: bytes) -> None:
    filas = _tabla_con(docx_bytes, "Fuente", "Tipo", "Cuota del sector")
    assert filas[1][0] == "wikipedia.org" and filas[1][6] == "Media"
    assert filas[1][5].endswith("p. p.")


def test_exactitud_y_mitigaciones(docx_bytes: bytes, diferencial: DatosDiferencial) -> None:
    texto = _seccion(_texto(docx_bytes), 6)
    assert diferencial.exactitud is not None
    h = diferencial.exactitud[0]
    resumen = {f[0]: f[1] for f in _tabla_con(docx_bytes, "Afirmaciones", "Número")[1:]}
    assert resumen["Incorrectas"] == str(h.n_incorrectas) == "0"
    assert resumen["A revisar por la clínica"] == str(h.n_riesgo_regulatorio)
    assert "Ninguna afirmación incorrecta" in texto
    afirmaciones = _tabla_con(docx_bytes, "Motor", "Lo que dice el motor")
    assert 1 < len(afirmaciones) <= 13
    assert all(f[2] == "A revisar por la clínica" for f in afirmaciones[1:])
    # Sin auditoría web en el estado: ni título ni fallo.
    assert "Auditoría técnica de la web" not in texto


def test_escenarios_como_proyecciones(docx_bytes: bytes) -> None:
    texto = _seccion(_texto(docx_bytes), 9)
    assert ESCENARIOS_INTRO in texto and "no una predicción" in texto
    filas = _tabla_con(docx_bytes, "Escenario")
    completo = next(f for f in filas[1:] if f[0].startswith("Completo"))
    assert completo[0] == "Completo (el de esta propuesta)"
    assert completo[5] == "30,8 [27,8–34,0]"  # cifra canónica: 25,9 → 30,8
    assert len(completo[1].split(", ")) == 9  # 9 acciones
    assert "Score de Citación actual: 25,9." in texto


def test_seguimiento_con_metodo(docx_bytes: bytes) -> None:
    texto = _seccion(_texto(docx_bytes), 10)
    for titulo, _ in PILARES_SEGUIMIENTO:
        assert titulo in texto
    linea_base = {f[0]: f[1] for f in _tabla_con(docx_bytes, "Indicador", "Valor")[1:]}
    assert linea_base["Score de Citación"] == "25,9 [21,6–30,1]"
    assert linea_base["Posición frente a la competencia"] == "8.ª de 10"
    assert ["Cadencia", "Trimestral"] in _tabla_con(docx_bytes, "Metodología")
    # Primera medición: no hay deltas que contar.
    assert "Cambios fuera del margen frente a la medición anterior" not in texto


def test_sin_diferencial_lo_dice_y_no_falla(proposal: Proposal) -> None:
    texto = _texto(DocxExporter().render(proposal, formato="docx"))
    posiciones = [texto.index(s) for s in SECCIONES]
    assert posiciones == sorted(posiciones)
    assert "{{" not in texto and "{%" not in texto
    assert texto.count(NO_CALCULADO) >= 5
    assert "Seguimiento con método" in texto


def test_del_estado_tolera_claves_ausentes_o_invalidas(diferencial: DatosDiferencial) -> None:
    estado = {
        "rigor": diferencial.rigor.model_dump() if diferencial.rigor else None,
        "auditoria_web": {"forma": "desconocida"},
        "otra_clave": 1,
    }
    datos = DatosDiferencial.del_estado(estado)
    assert datos.rigor == diferencial.rigor
    assert datos.auditoria_web is None and datos.competidores is None


def _auditoria_alborada() -> AuditoriaWeb:
    dominio = "alboradavision.example"
    fetcher = FixtureWebFetcher()
    respuestas = {
        r.id: fetcher.get(f"https://{dominio}{r.ruta}", dominio=dominio, accept=r.accept)
        for r in RECURSOS
    }
    return analizar_web(
        respuestas,
        study_id=STUDY_ID,
        brand_id="clinica-alborada-vision",
        dominio=dominio,
        origen="fixture",
        fetcher=fetcher.name,
        obtenido_en=datetime(2026, 10, 21),
    )


def test_auditoria_web_solo_si_el_estado_la_trae(
    proposal: Proposal, diferencial: DatosDiferencial
) -> None:
    con = diferencial.model_copy(update={"auditoria_web": _auditoria_alborada()})
    contenido = DocxExporter().render(proposal, formato="docx", diferencial=con)
    texto = _seccion(_texto(contenido), 6)
    assert "Auditoría técnica de la web: alboradavision.example" in texto
    assert AUDITORIA_FIXTURE in texto
    bots = {f[0]: f[2] for f in _tabla_con(contenido, "Bot", "Motor")[1:]}
    assert bots["GPTBot"] == "No"
    assert "Comprobaciones que fallan" in texto
