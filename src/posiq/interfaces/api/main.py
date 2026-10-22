"""API de posiq.: `uvicorn posiq.interfaces.api.main:app`.

- `POST /studies` crea el estudio y lanza el grafo en segundo plano → `{study_id}`.
- `GET /studies/{id}` estado, revisión pendiente, coste, errores y lo producido hasta ahora.
- `POST /studies/{id}/resume` reanuda una revisión (`competidores` o `plan`) → 202.
- `GET /studies/{id}/remeasure/estimate` coste de re-medir (mismo banco, sin intake ni mercado).
- `POST /studies/{id}/remeasure` re-mide un estudio completado (mismo brief y mercado) → 201. Si
  cuesta algo, exige `{confirmado: true, coste_estimado_eur}` y 409 si la estimación lo supera.
  Mide con la configuración VIGENTE de la marca y guarda las versiones usadas; 409 si ya hay una
  re-medición en curso del mismo estudio.
- `GET /brands` marcas del usuario con sus mediciones; `GET /brands/{id}/serie` su serie temporal.
- `GET|POST /studies/{id}/scenarios` escenarios del plan (preajustes y del consultor).
- `POST /studies/{id}/scenarios/{eid}/elegir` el escenario que usará la propuesta.
- `GET|PUT /brands/{id}/perfil` y `.../mercado` (`?historial=1`): configuración versionada con
  `vigente_desde` y `aviso_impacto` (`application/configuracion.py`).
- `PATCH /studies/{id}/acciones/{action_id} {estado, nota?}` estado de una acción del plan.
- `GET /studies/{id}/export?formato=docx|pdf` Word o PDF (409 sin propuesta; PDF sin soffice, 503).
- `GET /studies/{id}/evaluations?ids=a,b` evidencias: evaluación con motor, prompt y respuesta.
- `GET /studies/{id}/traces` trazas del estudio (modelo, tokens, coste por llamada).
- `GET /estimate?modo=&n_prompts=&n_motores=&repeticiones=` coste estimado antes de ejecutar.
- `GET /studies?brand_id=` lista de estudios.
- `GET /health`.
- `/auth/*` registro, verificación por código, login, logout y `me` (ver `auth.py`).

Todo salvo `/health`, `/auth/register|verify|verify/reenviar|login` (y la documentación
`/docs`) exige `Authorization: Bearer <token>` (401). Cada estudio es del usuario que lo crea: los
ajenos responden 404, como si no existieran (ADR 0015).

`POST /studies` acepta JSON `{brief_texto, modo}`, un fichero subido como cuerpo (`text/markdown`,
`text/plain` o `.docx`, con `?modo=`) o un formulario multipart con `fichero` y `modo` (este último
requiere `python-multipart`).
"""

import io
from collections.abc import AsyncIterator
from contextlib import asynccontextmanager
from pathlib import Path
from typing import Annotated, Any, Literal

from fastapi import BackgroundTasks, Body, Depends, FastAPI, HTTPException, Query, Request, status
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import Response
from pydantic import ValidationError

from posiq.adapters.export import (
    DatosDiferencial,
    DocxExporter,
    PdfExporter,
    SofficeNotFoundError,
)
from posiq.adapters.repo.accounts import SqlAccountRepo
from posiq.adapters.repo.tracking import SqlTrackingRepo
from posiq.application import seguimiento
from posiq.application.accounts import AccountService, Usuario, now_utc
from posiq.application.configuracion import BrandNotFound, RemeasureConflict, TrackingService
from posiq.application.deps import db
from posiq.application.diferencial import EscenarioGuardado
from posiq.application.estimate import (
    KIT_MOTORES,
    KIT_PROMPTS,
    KIT_REPETICIONES,
    CostEstimate,
    estimate_study,
)
from posiq.application.kit import load_sources
from posiq.application.runner import RemeasureError, ResumeError, ScenarioError, StudyRunner
from posiq.application.tracking import (
    AccionActualizada,
    AccionBody,
    EstadoAccionGuardado,
    MercadoBody,
    PerfilBody,
    VersionCreada,
    VersionMercadoMarca,
    VersionPerfilMarca,
)
from posiq.domain.scenario_schemas import EntradaEscenario
from posiq.domain.schemas import ModoEstudio, Study, Trace
from posiq.domain.tracking_schemas import SerieMarca
from posiq.interfaces.api.auth import CurrentUser, require_json
from posiq.interfaces.api.auth import router as auth_router
from posiq.interfaces.api.schemas import (
    AVISO_KIT,
    BrandSummary,
    Evidence,
    RemeasureCreated,
    RemeasureRequest,
    ResumeAccepted,
    ResumeBody,
    StudyCreate,
    StudyCreated,
    StudyView,
)
from posiq.interfaces.api.views import evidence
from posiq.interfaces.api.wiring import runner_from_settings
from posiq.settings import Settings

DOCX = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
TEXT_TYPES = ("text/markdown", "text/plain", "text/x-markdown")
LOCALHOST = r"https?://(localhost|127\.0\.0\.1)(:\d+)?"
# Margen de redondeo al comparar la estimación actual con la confirmada (céntimos).
TOLERANCIA_EUR = 0.005
MARGEN_MULTIPART = 64 * 1024


def docx_text(data: bytes) -> str:
    from docx import Document

    doc = Document(io.BytesIO(data))
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text for cell in row.cells))
    return "\n".join(p for p in parts if p.strip())


def file_text(nombre: str, content_type: str, data: bytes) -> str:
    if content_type.startswith(DOCX) or nombre.lower().endswith(".docx"):
        return docx_text(data)
    if content_type.startswith(TEXT_TYPES) or nombre.lower().endswith((".md", ".txt")):
        return data.decode("utf-8-sig")
    raise HTTPException(
        status.HTTP_415_UNSUPPORTED_MEDIA_TYPE, f"formato no admitido: {content_type or nombre}"
    )


def too_large(limite: int) -> HTTPException:
    return HTTPException(
        status.HTTP_413_CONTENT_TOO_LARGE, f"el brief supera el máximo de {limite} bytes"
    )


async def read_capped(request: Request, maximo: int, limite: int) -> bytes:
    """Cuerpo leído en streaming; 413 en cuanto pasa de `maximo` bytes.

    No se fía de `Content-Length` (puede faltar, p. ej. con `Transfer-Encoding: chunked`): corta
    al leer, sin cargar en memoria más de `maximo` bytes más un trozo. El cuerpo leído queda en
    `request._body`, que es lo que Starlette reutiliza en `request.body()` y `request.form()`.
    """
    partes: list[bytes] = []
    total = 0
    async for trozo in request.stream():
        total += len(trozo)
        if total > maximo:
            raise too_large(limite)
        partes.append(trozo)
    data = b"".join(partes)
    request._body = data
    return data


async def read_brief(request: Request, modo: ModoEstudio | None, limite: int) -> StudyCreate:
    content_type = request.headers.get("content-type", "")
    declarado = request.headers.get("content-length", "")
    multipart = content_type.startswith("multipart/form-data")
    # Margen para la cabecera multipart; el tamaño exacto se comprueba al leer.
    maximo = limite + MARGEN_MULTIPART if multipart else limite
    if declarado.isdigit() and int(declarado) > maximo:
        raise too_large(limite)
    data = await read_capped(request, maximo, limite)
    try:
        if multipart:
            body = await read_form(request, modo, limite)
        elif content_type.startswith("application/json"):
            body = StudyCreate.model_validate_json(data)
        else:
            nombre = request.query_params.get("nombre", "")
            texto = file_text(nombre, content_type, data)
            body = StudyCreate.model_validate({"brief_texto": texto, "modo": modo or "kit"})
    except ValidationError as exc:
        raise HTTPException(
            status.HTTP_422_UNPROCESSABLE_CONTENT, exc.errors(include_url=False)
        ) from exc
    if len(body.brief_texto.encode("utf-8")) > limite:
        raise too_large(limite)
    return body


async def read_form(request: Request, modo: ModoEstudio | None, limite: int) -> StudyCreate:
    try:
        form = await request.form()
    except AssertionError as exc:  # Starlette sin python-multipart
        raise HTTPException(
            status.HTTP_415_UNSUPPORTED_MEDIA_TYPE,
            "multipart no disponible en este servidor; envía JSON o el fichero como "
            "cuerpo (Content-Type text/markdown o .docx) con ?modo=",
        ) from exc
    fichero = form.get("fichero")
    texto = form.get("brief_texto")
    if fichero is not None and not isinstance(fichero, str):
        data = await fichero.read()
        if len(data) > limite:
            raise too_large(limite)
        texto = file_text(fichero.filename or "", fichero.content_type or "", data)
    return StudyCreate.model_validate(
        {"brief_texto": texto, "modo": form.get("modo") or modo or "kit"}
    )


def accounts_for(runner: StudyRunner, settings: Settings) -> AccountService:
    """Cuentas en la misma base que los estudios (el repositorio SQL expone su `engine`)."""
    engine = getattr(runner.repo, "engine", None)
    if engine is None:
        raise TypeError("las cuentas necesitan un repositorio SQL con `engine`")
    return AccountService(SqlAccountRepo(engine), settings)


def tracking_for(runner: StudyRunner) -> TrackingService:
    """Configuración versionada y estado de acciones en la misma base que los estudios."""
    engine = getattr(runner.repo, "engine", None)
    if engine is None:
        raise TypeError("el seguimiento necesita un repositorio SQL con `engine`")
    return TrackingService(SqlTrackingRepo(engine), runner)


def create_app(
    runner: StudyRunner | None = None,
    settings: Settings | None = None,
    accounts: AccountService | None = None,
) -> FastAPI:
    """App con `runner` inyectado (tests) o construido desde `settings` al arrancar.

    Con `runner` inyectado, las cuentas se montan ya sobre su repositorio (sin esperar al
    `lifespan`, que `httpx.ASGITransport` no ejecuta)."""

    @asynccontextmanager
    async def lifespan(app: FastAPI) -> AsyncIterator[None]:
        if runner is not None:
            yield
            return
        app.state.settings = settings or Settings()
        async with runner_from_settings(app.state.settings) as built:
            app.state.runner = built
            app.state.accounts = accounts or accounts_for(built, app.state.settings)
            app.state.tracking = tracking_for(built)
            yield

    app = FastAPI(title="posiq.", lifespan=lifespan)
    app.state.runner = runner
    app.state.settings = settings
    app.state.accounts = accounts
    if runner is not None and accounts is None:
        app.state.accounts = accounts_for(runner, settings or Settings())
    app.state.tracking = tracking_for(runner) if runner is not None else None
    app.include_router(auth_router)
    app.add_middleware(
        CORSMiddleware,
        allow_origin_regex=LOCALHOST,
        allow_methods=["*"],
        allow_headers=["*"],
    )

    def get_runner(request: Request) -> StudyRunner:
        current: StudyRunner | None = request.app.state.runner
        if current is None:
            raise HTTPException(status.HTTP_503_SERVICE_UNAVAILABLE, "la API aún no ha arrancado")
        return current

    def get_accounts(request: Request) -> AccountService:
        current: AccountService = request.app.state.accounts
        return current

    def get_study(request: Request, r: StudyRunner, user: Usuario, study_id: str) -> Study:
        """El estudio si es de `user`; 404 si no existe o es de otro (no se revela cuál)."""
        try:
            study: Study | None = r.repo.get_study(study_id)
        except KeyError:
            study = None
        if study is None or not get_accounts(request).can_access(user, study_id):
            raise HTTPException(status.HTTP_404_NOT_FOUND, f"no existe el estudio {study_id}")
        return study

    def get_settings(request: Request) -> Settings:
        current: Settings | None = request.app.state.settings
        return current or Settings()

    def get_tracking(request: Request) -> TrackingService:
        get_runner(request)
        current: TrackingService = request.app.state.tracking
        return current

    def brand_404(brand_id: str) -> HTTPException:
        return HTTPException(status.HTTP_404_NOT_FOUND, f"no existe la marca {brand_id}")

    @app.get("/health")
    async def health() -> dict[str, str]:
        return {"status": "ok"}

    @app.post("/studies", status_code=status.HTTP_201_CREATED)
    async def create_study(
        request: Request,
        background: BackgroundTasks,
        user: CurrentUser,
        modo: Annotated[ModoEstudio | None, Query()] = None,
    ) -> StudyCreated:
        r = get_runner(request)
        body = await read_brief(request, modo, get_settings(request).brief_max_bytes)
        try:
            study = r.create(body.modo)
        except ValueError as exc:
            raise HTTPException(status.HTTP_422_UNPROCESSABLE_CONTENT, str(exc)) from exc
        get_accounts(request).repo.set_owner(study.id, user.id)
        background.add_task(r.start, study, body.brief_texto)
        return StudyCreated(study_id=study.id, aviso=AVISO_KIT if body.modo == "kit" else None)

    @app.get("/studies")
    async def list_studies(
        request: Request, user: CurrentUser, brand_id: str | None = None
    ) -> list[Study]:
        studies = get_runner(request).repo.list_studies(brand_id)
        visibles = set(get_accounts(request).visible_ids(user, [s.id for s in studies]))
        return [s for s in studies if s.id in visibles]

    @app.get("/studies/{study_id}")
    async def read_study(request: Request, study_id: str, user: CurrentUser) -> StudyView:
        r = get_runner(request)
        study = get_study(request, r, user, study_id)
        values, _ = await r.state(study_id)
        errores = list(values.get("errores", []))
        if study_id in r.fallos:
            errores.append(r.fallos[study_id])
        # El grafo guarda "none" en su estado; la web recibe null.
        pendiente = values.get("revision_pendiente", "none")
        return StudyView(
            study=study,
            estado=r.estado(study),  # type: ignore[arg-type]
            revision_pendiente=None if pendiente == "none" else pendiente,
            coste_acumulado_eur=values.get("coste_acumulado_eur", 0.0),
            errores=errores,
            brief=values.get("brief"),
            market=values.get("market"),
            metrics=values.get("metrics"),
            metrics_delta=values.get("metrics_delta"),
            findings=values.get("findings"),
            plan=values.get("plan"),
            simulacion=values.get("simulacion"),
            calidad_plan=values.get("calidad_plan"),
            rigor=values.get("rigor"),
            competidores=values.get("competidores"),
            escenarios=values.get("escenarios"),
            oportunidades_fuente=values.get("oportunidades_fuente"),
            exactitud=values.get("exactitud"),
            mitigaciones=values.get("mitigaciones"),
            auditoria_web=values.get("auditoria_web"),
            hallazgos_web=values.get("hallazgos_web"),
            proposal=values.get("proposal"),
            datos_sinteticos=study.modo == "kit",
            acciones_estado=await db(get_tracking(request).repo.action_statuses, study_id),
            configuracion=await db(get_tracking(request).repo.study_config, study_id),
        )

    @app.post(
        "/studies/{study_id}/resume",
        status_code=status.HTTP_202_ACCEPTED,
        dependencies=[Depends(require_json)],
    )
    async def resume_study(
        request: Request,
        study_id: str,
        body: ResumeBody,
        background: BackgroundTasks,
        user: CurrentUser,
    ) -> ResumeAccepted:
        r = get_runner(request)
        get_study(request, r, user, study_id)
        try:
            await r.check_resume(study_id, body.revision, body.payload)
        except ResumeError as exc:
            code = (
                status.HTTP_409_CONFLICT if exc.conflict else status.HTTP_422_UNPROCESSABLE_CONTENT
            )
            raise HTTPException(code, str(exc)) from exc
        r.en_marcha.add(study_id)  # evita un segundo resume antes de que arranque el primero
        background.add_task(r.resume, study_id, body.payload)
        return ResumeAccepted(study_id=study_id, revision=body.revision)

    async def remeasure_estimate(request: Request, r: StudyRunner, study: Study) -> CostEstimate:
        """Coste de re-medir `study`: reutiliza su banco de prompts; sin intake ni mercado."""
        values, _ = await r.state(study.id)
        n_prompts = len(values.get("prompts") or []) or KIT_PROMPTS
        return estimate_study(
            study.modo, get_settings(request), n_prompts=n_prompts, remedicion=True
        )

    @app.get("/studies/{study_id}/remeasure/estimate")
    async def read_remeasure_estimate(
        request: Request, study_id: str, user: CurrentUser
    ) -> CostEstimate:
        r = get_runner(request)
        return await remeasure_estimate(request, r, get_study(request, r, user, study_id))

    @app.post("/studies/{study_id}/remeasure", status_code=status.HTTP_201_CREATED)
    async def remeasure_study(
        request: Request,
        study_id: str,
        background: BackgroundTasks,
        user: CurrentUser,
        confirmacion: Annotated[RemeasureRequest | None, Body()] = None,
    ) -> RemeasureCreated:
        r = get_runner(request)
        original = get_study(request, r, user, study_id)
        coste = (await remeasure_estimate(request, r, original)).coste_estimado_eur
        if coste > 0 and confirmacion is None:
            raise HTTPException(
                status.HTTP_409_CONFLICT,
                f"la re-medición cuesta {coste:.2f} €: confirma con "
                '{"confirmado": true, "coste_estimado_eur": …}',
            )
        if confirmacion is not None and coste > confirmacion.coste_estimado_eur + TOLERANCIA_EUR:
            raise HTTPException(
                status.HTTP_409_CONFLICT,
                f"la estimación ha subido a {coste:.2f} € (confirmaste "
                f"{confirmacion.coste_estimado_eur:.2f} €): vuelve a confirmar",
            )
        try:
            study = await get_tracking(request).remeasure(
                user.id, visible_studies(request, r, user), original
            )
        except RemeasureError as exc:
            raise HTTPException(status.HTTP_409_CONFLICT, str(exc)) from exc
        # Hereda el propietario del original (o el usuario demo, si el original no tenía).
        accounts = get_accounts(request)
        accounts.repo.set_owner(study.id, accounts.repo.owner_of(study_id) or user.id)
        r.en_marcha.add(study.id)  # en_curso desde ya, antes de que arranque la tarea
        background.add_task(r.continue_, study)
        return RemeasureCreated(study_id=study.id, remedicion_de=study_id)

    def scenario_http_error(exc: ScenarioError) -> HTTPException:
        code = status.HTTP_409_CONFLICT if exc.conflict else status.HTTP_422_UNPROCESSABLE_CONTENT
        return HTTPException(code, str(exc))

    @app.get("/studies/{study_id}/scenarios")
    async def list_scenarios(
        request: Request, study_id: str, user: CurrentUser
    ) -> list[EscenarioGuardado]:
        r = get_runner(request)
        get_study(request, r, user, study_id)
        return await r.scenarios(study_id)

    @app.post(
        "/studies/{study_id}/scenarios",
        status_code=status.HTTP_201_CREATED,
        dependencies=[Depends(require_json)],
    )
    async def create_scenario(
        request: Request, study_id: str, entrada: EntradaEscenario, user: CurrentUser
    ) -> EscenarioGuardado:
        r = get_runner(request)
        get_study(request, r, user, study_id)
        try:
            return await r.add_scenario(study_id, entrada)
        except ScenarioError as exc:
            raise scenario_http_error(exc) from exc

    @app.post("/studies/{study_id}/scenarios/{escenario_id}/elegir")
    async def choose_scenario(
        request: Request, study_id: str, escenario_id: str, user: CurrentUser
    ) -> list[EscenarioGuardado]:
        r = get_runner(request)
        get_study(request, r, user, study_id)
        try:
            return await r.choose_scenario(study_id, escenario_id)
        except ScenarioError as exc:
            raise scenario_http_error(exc) from exc
        except KeyError:
            raise HTTPException(
                status.HTTP_404_NOT_FOUND, f"no existe el escenario {escenario_id}"
            ) from None

    @app.patch("/studies/{study_id}/acciones/{action_id}")
    async def update_action(
        request: Request, study_id: str, action_id: str, body: AccionBody, user: CurrentUser
    ) -> AccionActualizada:
        r = get_runner(request)
        get_study(request, r, user, study_id)
        values, _ = await r.state(study_id)
        plan = values.get("plan")
        if plan is None or action_id not in {a.id for a in plan.acciones}:
            raise HTTPException(
                status.HTTP_404_NOT_FOUND, f"no existe la acción {action_id} en {study_id}"
            )
        guardado = EstadoAccionGuardado(
            estado=body.estado, nota=body.nota, actualizado_en=now_utc()
        )
        repo = get_tracking(request).repo
        await db(repo.add_action_status, study_id, action_id, user.id, guardado)
        return AccionActualizada(action_id=action_id, **guardado.model_dump())

    @app.get("/studies/{study_id}/evaluations")
    async def read_evaluations(
        request: Request, study_id: str, user: CurrentUser, ids: str = ""
    ) -> list[Evidence]:
        r = get_runner(request)
        get_study(request, r, user, study_id)
        values, _ = await r.state(study_id)
        return evidence(values, [i for i in ids.split(",") if i])

    @app.get("/studies/{study_id}/traces")
    async def read_traces(request: Request, study_id: str, user: CurrentUser) -> list[Trace]:
        r = get_runner(request)
        get_study(request, r, user, study_id)
        values, _ = await r.state(study_id)
        return list(values.get("traces", []))

    def visible_studies(request: Request, r: StudyRunner, user: Usuario) -> list[Study]:
        studies = r.repo.list_studies(None)
        visibles = set(get_accounts(request).visible_ids(user, [s.id for s in studies]))
        return [s for s in studies if s.id in visibles]

    @app.get("/brands")
    async def list_brands(request: Request, user: CurrentUser) -> list[BrandSummary]:
        studies = visible_studies(request, get_runner(request), user)
        return [
            BrandSummary(
                brand_id=m.brand_id,
                nombre=m.nombre,
                n_mediciones=m.n_mediciones,
                ultima_medicion=m.ultima_medicion,
            )
            for m in seguimiento.resumen_marcas(studies)
        ]

    @app.get("/brands/{brand_id}/serie")
    async def brand_series(request: Request, brand_id: str, user: CurrentUser) -> SerieMarca:
        """Serie con método (ADR 0016) de las mediciones completadas de la marca del usuario."""
        r = get_runner(request)
        propios = [s for s in visible_studies(request, r, user) if s.brand_id == brand_id]
        catalogo = load_sources(str(Path(get_settings(request).kit_dir).resolve()))
        meds = await seguimiento.mediciones(r, propios, catalogo)
        try:
            serie = await db(seguimiento.serie, brand_id, meds)
        except ValueError as exc:
            raise HTTPException(status.HTTP_422_UNPROCESSABLE_CONTENT, str(exc)) from exc
        if serie is None:
            raise HTTPException(
                status.HTTP_404_NOT_FOUND, f"no hay mediciones completadas de la marca {brand_id}"
            )
        return serie

    # --- Configuración versionada por marca (ADR 0016) ----------------------------------------

    @app.get("/brands/{brand_id}/perfil", response_model=None)
    async def brand_profile(
        request: Request, brand_id: str, user: CurrentUser, historial: bool = False
    ) -> VersionPerfilMarca | list[VersionPerfilMarca]:
        try:
            versiones = await get_tracking(request).perfiles(
                user.id, visible_studies(request, get_runner(request), user), brand_id
            )
        except BrandNotFound as exc:
            raise brand_404(brand_id) from exc
        if historial:
            return versiones
        if not versiones:
            raise HTTPException(
                status.HTTP_404_NOT_FOUND, f"la marca {brand_id} aún no tiene perfil"
            )
        return versiones[-1]

    @app.put("/brands/{brand_id}/perfil", status_code=status.HTTP_201_CREATED)
    async def put_brand_profile(
        request: Request, brand_id: str, body: PerfilBody, user: CurrentUser
    ) -> VersionCreada:
        try:
            return await get_tracking(request).guardar_perfil(
                user.id, visible_studies(request, get_runner(request), user), brand_id, body.perfil
            )
        except BrandNotFound as exc:
            raise brand_404(brand_id) from exc
        except RemeasureConflict as exc:
            raise HTTPException(status.HTTP_409_CONFLICT, str(exc)) from exc

    @app.get("/brands/{brand_id}/mercado", response_model=None)
    async def brand_market(
        request: Request, brand_id: str, user: CurrentUser, historial: bool = False
    ) -> VersionMercadoMarca | list[VersionMercadoMarca]:
        try:
            versiones = await get_tracking(request).mercados(
                user.id, visible_studies(request, get_runner(request), user), brand_id
            )
        except BrandNotFound as exc:
            raise brand_404(brand_id) from exc
        if historial:
            return versiones
        if not versiones:
            raise HTTPException(
                status.HTTP_404_NOT_FOUND, f"la marca {brand_id} aún no tiene mercado confirmado"
            )
        return versiones[-1]

    @app.put("/brands/{brand_id}/mercado", status_code=status.HTTP_201_CREATED)
    async def put_brand_market(
        request: Request, brand_id: str, body: MercadoBody, user: CurrentUser
    ) -> VersionCreada:
        try:
            return await get_tracking(request).guardar_mercado(
                user.id, visible_studies(request, get_runner(request), user), brand_id, body.marcas
            )
        except BrandNotFound as exc:
            raise brand_404(brand_id) from exc
        except RemeasureConflict as exc:
            raise HTTPException(status.HTTP_409_CONFLICT, str(exc)) from exc
        except ValueError as exc:
            raise HTTPException(status.HTTP_422_UNPROCESSABLE_CONTENT, str(exc)) from exc

    @app.get("/estimate")
    async def estimate(
        request: Request,
        user: CurrentUser,
        modo: ModoEstudio,
        n_prompts: Annotated[int, Query(ge=1)] = KIT_PROMPTS,
        n_motores: Annotated[int, Query(ge=1)] = KIT_MOTORES,
        repeticiones: Annotated[int, Query(ge=1)] = KIT_REPETICIONES,
    ) -> CostEstimate:
        return estimate_study(
            modo,
            get_settings(request),
            n_prompts=n_prompts,
            n_motores=n_motores,
            repeticiones=repeticiones,
        )

    @app.get("/studies/{study_id}/export")
    async def export_study(
        request: Request,
        study_id: str,
        user: CurrentUser,
        formato: Literal["docx", "pdf"] = "docx",
    ) -> Any:
        r = get_runner(request)
        get_study(request, r, user, study_id)
        values, _ = await r.state(study_id)
        proposal = values.get("proposal")
        if proposal is None:
            raise HTTPException(status.HTTP_409_CONFLICT, "la propuesta aún no está completada")
        exporter = DocxExporter() if formato == "docx" else PdfExporter()
        # Rigor, competidores, escenarios...: cifras del estado que no viven en la `Proposal`.
        diferencial = DatosDiferencial.del_estado(values)
        try:
            data = await db(exporter.render, proposal, formato=formato, diferencial=diferencial)
        except SofficeNotFoundError as exc:
            raise HTTPException(status.HTTP_503_SERVICE_UNAVAILABLE, str(exc)) from exc
        return Response(
            data,
            media_type=DOCX if formato == "docx" else "application/pdf",
            headers={"Content-Disposition": f'attachment; filename="posiq-{study_id}.{formato}"'},
        )

    return app


app = create_app()
