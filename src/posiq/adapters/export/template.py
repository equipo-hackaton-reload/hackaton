"""Plantilla docxtpl de la propuesta, construida con python-docx (reproducible, sin edición a mano).

La plantilla solo coloca valores ya formateados por `docx.build_context`: no hace aritmética.
Las diez secciones del formato del entregable (doc 06) más «Lo que nos habéis pedido» (2.ª): el
encargo literal del brief y dónde se trata cada punto (ADR 0013), «Exactitud y lectura técnica»
(7.ª) y «Escenarios y simulación» (10.ª). «Medición y seguimiento» pasa a ser la oferta del
servicio: «Seguimiento con método» (11.ª). El rigor (intervalos) va en el diagnóstico y en cada
hallazgo; competidores declarados frente a los de la IA y oportunidades de fuente, en la
competencia. La auditoría técnica de la web solo aparece si el estudio la trae. Se regenera con
`uv run python scripts/build_docx_template.py`.
"""

from pathlib import Path

from docx import Document
from docx.document import Document as DocumentObject
from docx.enum.section import WD_ORIENT
from docx.shared import Pt

TEMPLATE_PATH = Path(__file__).parent / "templates" / "propuesta.docx"

SECCIONES: tuple[str, ...] = (
    "1. Resumen ejecutivo",
    "2. Lo que nos habéis pedido",
    "3. Contexto y alcance",
    "4. Diagnóstico de visibilidad",
    "5. Análisis de la competencia",
    "6. Hallazgos clave",
    "7. Exactitud y lectura técnica de la marca",
    "8. Estrategia recomendada",
    "9. Plan de acción por fases",
    "10. Escenarios y simulación del plan",
    "11. Seguimiento con método",
    "12. Riesgos y supuestos",
    "13. Hoja de ruta y presupuesto",
)


def _p(doc: DocumentObject, text: str, style: str | None = None) -> None:
    doc.add_paragraph(text, style=style)


def _tag(doc: DocumentObject, stmt: str) -> None:
    """Párrafo de control de docxtpl (`{%p ... %}`), que desaparece al renderizar."""
    doc.add_paragraph(f"{{%p {stmt} %}}")


def _texto(doc: DocumentObject, expr: str) -> None:
    """Lista de párrafos de texto libre."""
    _tag(doc, f"for par in {expr}")
    _p(doc, "{{ par }}")
    _tag(doc, "endfor")


def _vinetas(doc: DocumentObject, expr: str) -> None:
    _tag(doc, f"for item in {expr}")
    _p(doc, "{{ item }}", "List Bullet")
    _tag(doc, "endfor")


def _tabla(doc: DocumentObject, cabeceras: list[str], expr: str, campos: list[str]) -> None:
    """Tabla con cabecera y una fila por elemento de `expr` (bucle `{%tr %}` de docxtpl)."""
    tabla = doc.add_table(rows=4, cols=len(cabeceras))
    tabla.style = "Light Grid Accent 1"
    for celda, texto in zip(tabla.rows[0].cells, cabeceras, strict=True):
        celda.text = texto
    tabla.rows[1].cells[0].text = f"{{%tr for r in {expr} %}}"
    for celda, campo in zip(tabla.rows[2].cells, campos, strict=True):
        celda.text = f"{{{{ r.{campo} }}}}"
    tabla.rows[3].cells[0].text = "{%tr endfor %}"
    doc.add_paragraph()


def _ficha(doc: DocumentObject, filas: list[tuple[str, str]]) -> None:
    """Tabla de dos columnas campo → valor."""
    tabla = doc.add_table(rows=len(filas), cols=2)
    tabla.style = "Light List Accent 1"
    for fila, (campo, valor) in zip(tabla.rows, filas, strict=True):
        fila.cells[0].text = campo
        fila.cells[1].text = valor
    doc.add_paragraph()


def _si_hay(doc: DocumentObject, bloque: str) -> None:
    """Abre un bloque que solo se pinta si el estudio lo calculó (`<bloque>.hay`)."""
    _tag(doc, f"if {bloque}.hay")


def _si_no_hay(doc: DocumentObject, bloque: str) -> None:
    """Cierra el bloque: si no se calculó, el documento lo dice (`<bloque>.aviso`)."""
    _tag(doc, "else")
    _p(doc, f"{{{{ {bloque}.aviso }}}}")
    _tag(doc, "endif")


def _cabecera_y_pie(doc: DocumentObject) -> None:
    section = doc.sections[0]
    section.orientation = WD_ORIENT.PORTRAIT
    section.header.paragraphs[0].text = "{{r cabecera }}"
    section.footer.paragraphs[0].text = "{{ pie }}"


def _portada(doc: DocumentObject) -> None:
    _p(doc, "{{ logo }}")
    doc.add_heading("{{ titulo }}", level=0)
    _p(doc, "{{ marca }} · {{ sector }} · {{ fecha }}")
    _tag(doc, "if rotulo_sintetico")
    _p(doc, "{{ rotulo_sintetico }}", "Intense Quote")
    _p(doc, "{{ aviso_sintetico }}")
    _tag(doc, "endif")


def _encargo(doc: DocumentObject) -> None:
    """Lo que pidió el cliente, tal cual, y la tabla de cobertura punto → dónde se trata."""
    doc.add_heading(SECCIONES[1], level=1)
    _p(doc, "{{ encargo.intro }}")
    _tag(doc, "for l in encargo.listas")
    doc.add_heading("{{ l.titulo }}", level=2)
    _vinetas(doc, "l.puntos")
    _tag(doc, "endfor")
    doc.add_heading("Dónde se trata cada punto", level=2)
    _tabla(
        doc,
        ["Punto", "Lo que pedisteis", "Dónde se trata", "Estado"],
        "encargo.cobertura",
        ["tipo", "texto", "donde", "estado"],
    )


def _diagnostico(doc: DocumentObject) -> None:
    doc.add_heading(SECCIONES[3], level=1)
    _texto(doc, "diag.texto")
    _ficha(
        doc,
        [
            ("Score de Citación global (0-100)", "{{ diag.score }}"),
            ("Aparición (máx. 30)", "{{ diag.aparicion }}"),
            ("Posición (máx. 20)", "{{ diag.posicion }}"),
            ("Fuerza de recomendación (máx. 30)", "{{ diag.fuerza }}"),
            ("Calidad de la fuente (máx. 20)", "{{ diag.calidad_fuente }}"),
            ("Evaluaciones", "{{ diag.n_evaluaciones }}"),
            ("Versión del score", "{{ diag.score_version }}"),
        ],
    )
    _p(doc, "{{ diag.nota_espontaneo }}")
    cabeceras = ["Segmento", "Visibilidad espontánea", "n", "Apariciones", "Fuerza media", "Líder"]
    campos = ["valor", "tasa", "n", "apariciones", "fuerza", "lider"]
    for titulo, clave in (
        ("Visibilidad espontánea por motor", "motor"),
        ("Visibilidad espontánea por fase del embudo", "fase"),
        ("Visibilidad espontánea por ciudad", "ciudad"),
        ("Visibilidad espontánea por perfil", "perfil"),
    ):
        doc.add_heading(titulo, level=2)
        _tabla(doc, cabeceras, f"diag.{clave}", campos)
    doc.add_heading("Visibilidad asistida (referencia, no comparable)", level=2)
    _tabla(doc, ["Segmento", "Visibilidad asistida", *cabeceras[2:]], "diag.asistido", campos)
    doc.add_heading("Calidad de las menciones", level=2)
    _tabla(
        doc,
        ["Tipo", "Apariciones", "Fuerza media", "Por tipo de mención", "Negativas", "Desconoc."],
        "diag.calidad",
        ["tipo", "apariciones", "fuerza", "por_tipo", "negativas", "desconocimiento"],
    )
    doc.add_heading("Fuentes que acompañan a la marca", level=2)
    _tabla(
        doc,
        ["Dominio", "Tipo", "Calidad", "Apariciones"],
        "diag.fuentes",
        ["dominio", "tipo", "calidad", "apariciones"],
    )
    _rigor(doc)


def _rigor(doc: DocumentObject) -> None:
    doc.add_heading("Margen de error de las cifras", level=2)
    _si_hay(doc, "rigor")
    _p(doc, "{{ rigor.intro }}")
    _ficha(
        doc,
        [
            ("Score de Citación [IC 95 %]", "{{ rigor.score }}"),
            ("Margen", "{{ rigor.margen }}"),
            ("Evaluaciones", "{{ rigor.n }}"),
        ],
    )
    doc.add_heading("Aparición espontánea por motor, con su margen", level=3)
    _tabla(
        doc,
        ["Motor", "Aparición [IC 95 %]", "Margen", "Respuestas", "Score del motor [IC 95 %]"],
        "rigor.motores",
        ["motor", "aparicion", "margen", "n", "score"],
    )
    _p(doc, "{{ rigor.metodo }}")
    _si_no_hay(doc, "rigor")


def _competencia(doc: DocumentObject) -> None:
    doc.add_heading(SECCIONES[4], level=1)
    _texto(doc, "comp.texto")
    doc.add_heading("Ranking por Score de Citación", level=2)
    _tabla(
        doc,
        ["#", "Marca", "Score", "Evaluaciones"],
        "comp.ranking",
        ["pos", "marca", "score", "n"],
    )
    doc.add_heading("Brecha de fuentes frente al líder", level=2)
    _tabla(
        doc,
        ["Dominio", "Tipo", "Con líder", "Con la marca", "Brecha", "Líder sin la marca"],
        "comp.brecha",
        ["dominio", "tipo", "con_lider", "con_marca", "brecha", "lider_sin_marca"],
    )
    cabeceras = [
        "Marca",
        "Posición",
        "Score",
        "Ventaja sobre la marca",
        "Visib. espontánea",
        "Nota",
    ]
    campos = ["marca", "puesto", "score", "ventaja", "tasa", "nota"]
    doc.add_heading("Competidores declarados frente a los que pone la IA", level=2)
    _si_hay(doc, "rivales")
    _p(doc, "{{ rivales.intro }}")
    _p(doc, "{{ rivales.foco }}")
    doc.add_heading("Los que declara el brief", level=3)
    _tabla(doc, cabeceras, "rivales.declarados", campos)
    doc.add_heading("Los que la IA pone por delante y el brief no declara", level=3)
    _tabla(doc, cabeceras, "rivales.por_delante", campos)
    _p(doc, "{{ rivales.umbral }}")
    _vinetas(doc, "rivales.notas")
    _si_no_hay(doc, "rivales")
    doc.add_heading("Mapa de fuentes y oportunidades", level=2)
    _si_hay(doc, "fuentes_op")
    _p(doc, "{{ fuentes_op.intro }}")
    _tabla(
        doc,
        ["Fuente", "Tipo", "Cuota del sector", "Con la marca", "Con competidores", "Brecha"]
        + ["Prioridad"],
        "fuentes_op.filas",
        ["fuente", "tipo", "cuota", "marca", "competidores", "brecha", "prioridad"],
    )
    _tag(doc, "if fuentes_op.resto")
    _p(doc, "{{ fuentes_op.resto }}")
    _tag(doc, "endif")
    _si_no_hay(doc, "fuentes_op")


def _hallazgos(doc: DocumentObject) -> None:
    doc.add_heading(SECCIONES[5], level=1)
    _tag(doc, "for h in hallazgos")
    doc.add_heading("{{ h.id }} · {{ h.titulo }}", level=2)
    _p(doc, "{{ h.descripcion }}")
    _ficha(
        doc,
        [
            ("Cifra", "{{ h.cifra }}"),
            ("Origen de la cifra", "{{ h.origen }}"),
            ("Fase del embudo", "{{ h.fase }}"),
            ("Impacto", "{{ h.impacto }}"),
            ("Confianza", "{{ h.confianza }}"),
            ("Evaluaciones que lo respaldan", "{{ h.n_evidencias }}"),
            ("Contraste estadístico", "{{ h.contraste }}"),
        ],
    )
    _tag(doc, "endfor")


def _plan(doc: DocumentObject) -> None:
    doc.add_heading(SECCIONES[8], level=1)
    _ficha(
        doc,
        [
            ("Presupuesto estimado", "{{ plan.presupuesto }}"),
            ("Respeta las restricciones del brief", "{{ plan.respeta }}"),
        ],
    )
    _tag(doc, "for f in plan.fases")
    doc.add_heading("{{ f.nombre }} · {{ f.semanas }} semanas", level=2)
    _p(doc, "Objetivo: {{ f.objetivo }}")
    _tag(doc, "for a in f.acciones")
    doc.add_heading("{{ a.id }} · {{ a.descripcion }}", level=3)
    _ficha(
        doc,
        [
            ("Palanca", "{{ a.palanca }}"),
            ("Fase del embudo", "{{ a.fase }}"),
            ("Motores objetivo", "{{ a.motores }}"),
            ("Esfuerzo / impacto", "{{ a.esfuerzo }} / {{ a.impacto }}"),
            ("Responsable", "{{ a.responsable }}"),
            ("Plazo", "{{ a.plazo }}"),
            ("KPI", "{{ a.kpi }}"),
            ("Entregables", "{{ a.entregables }}"),
            ("Dependencias", "{{ a.dependencias }}"),
            ("Hallazgos que la justifican", "{{ a.hallazgos }}"),
            ("Restricciones revisadas", "{{ a.restricciones }}"),
        ],
    )
    _tag(doc, "endfor")
    _tag(doc, "endfor")


def _exactitud(doc: DocumentObject) -> None:
    doc.add_heading(SECCIONES[6], level=1)
    doc.add_heading("Lo que dicen los motores de la marca", level=2)
    _si_hay(doc, "exactitud")
    _p(doc, "{{ exactitud.intro }}")
    _tabla(doc, ["Afirmaciones", "Número"], "exactitud.resumen", ["campo", "valor"])
    doc.add_heading("Afirmaciones a revisar", level=3)
    _tabla(
        doc,
        ["Motor", "Lo que dice el motor", "Estado", "Motivo"],
        "exactitud.revisar",
        ["motor", "texto", "estado", "motivo"],
    )
    _tag(doc, "if exactitud.resto")
    _p(doc, "{{ exactitud.resto }}")
    _tag(doc, "endif")
    doc.add_heading("Mitigaciones en origen", level=3)
    _tag(doc, "if exactitud.mitigaciones")
    _tabla(
        doc,
        ["Id", "Acción", "Fuente que lo publica", "Afirmaciones"],
        "exactitud.mitigaciones",
        ["id", "descripcion", "fuente", "afirmaciones"],
    )
    _tag(doc, "else")
    _p(doc, "{{ exactitud.sin_mitigaciones }}")
    _tag(doc, "endif")
    _si_no_hay(doc, "exactitud")
    # La auditoría técnica solo existe si el estudio la trae: sin ella no hay ni título.
    _tag(doc, "if auditoria.hay")
    doc.add_heading("Auditoría técnica de la web: {{ auditoria.dominio }}", level=2)
    _p(doc, "{{ auditoria.intro }}")
    _tag(doc, "if auditoria.aviso")
    _p(doc, "{{ auditoria.aviso }}", "Intense Quote")
    _tag(doc, "endif")
    _ficha(
        doc,
        [
            ("Preparación para los motores", "{{ auditoria.puntuacion }}"),
            ("Fecha de la auditoría", "{{ auditoria.fecha }}"),
        ],
    )
    _tabla(
        doc,
        ["Bloque", "Puntuación", "Correctas", "Fallan"],
        "auditoria.bloques",
        ["bloque", "puntuacion", "ok", "falla"],
    )
    doc.add_heading("Comprobaciones que fallan", level=3)
    _tabla(
        doc,
        ["Comprobación", "Detalle", "Motores afectados", "Evidencia"],
        "auditoria.fallos",
        ["titulo", "detalle", "motores", "evidencia"],
    )
    doc.add_heading("Acceso de los bots de IA", level=3)
    _tabla(doc, ["Bot", "Motor", "Permitido"], "auditoria.bots", ["bot", "motor", "permitido"])
    _tag(doc, "endif")


def _escenarios(doc: DocumentObject) -> None:
    doc.add_heading(SECCIONES[9], level=1)
    _si_hay(doc, "escenarios")
    _p(doc, "{{ escenarios.intro }}")
    _tag(doc, "if escenarios.hay_escenarios")
    doc.add_heading("Escenarios del plan", level=2)
    _tabla(
        doc,
        ["Escenario", "Acciones", "Aplazadas", "Hallazgos sin cubrir", "Esfuerzo"]
        + ["Score proyectado [mín.–máx.]"],
        "escenarios.escenarios",
        ["nombre", "mantenidas", "aplazadas", "sin_cubrir", "esfuerzo", "score"],
    )
    _tag(doc, "endif")
    _tag(doc, "if escenarios.hay_simulacion")
    doc.add_heading("Simulación del plan completo por fases", level=2)
    _p(doc, "Score de Citación actual: {{ escenarios.actual }}.")
    _tabla(
        doc,
        ["Al terminar", "Score proyectado [mín.–máx.]", "Visibilidad espontánea [mín.–máx.]"],
        "escenarios.fases",
        ["fase", "score", "tasa"],
    )
    doc.add_heading("Supuestos", level=3)
    _vinetas(doc, "escenarios.supuestos")
    _p(doc, "{{ escenarios.advertencia }}", "Intense Quote")
    _p(doc, "Modelo de simulación: {{ escenarios.modelo }}.")
    _tag(doc, "endif")
    _si_no_hay(doc, "escenarios")


def _seguimiento(doc: DocumentObject) -> None:
    doc.add_heading(SECCIONES[10], level=1)
    _p(doc, "{{ seguimiento.intro }}")
    _tag(doc, "for m in seguimiento.metodo")
    doc.add_heading("{{ m.titulo }}", level=3)
    _p(doc, "{{ m.texto }}")
    _tag(doc, "endfor")
    doc.add_heading("Línea base de este estudio", level=2)
    _tabla(doc, ["Indicador", "Valor"], "seguimiento.linea_base", ["campo", "valor"])
    _tag(doc, "if seguimiento.hay_remedicion")
    doc.add_heading("Cambios fuera del margen frente a la medición anterior", level=2)
    _tabla(
        doc,
        ["Indicador", "Cambio", "IC 95 % del cambio"],
        "seguimiento.deltas",
        ["indicador", "cambio", "ic"],
    )
    _tag(doc, "endif")
    doc.add_heading("Indicadores y cadencia de esta propuesta", level=2)
    _vinetas(doc, "medicion.kpis")
    _ficha(
        doc,
        [("Metodología", "{{ medicion.metodologia }}"), ("Cadencia", "{{ medicion.cadencia }}")],
    )


def build_template(path: Path = TEMPLATE_PATH) -> Path:
    """Escribe la plantilla en `path` y la devuelve."""
    doc = Document()
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"  # type: ignore[attr-defined]
    normal.font.size = Pt(10.5)  # type: ignore[attr-defined]
    _cabecera_y_pie(doc)
    _portada(doc)

    doc.add_heading(SECCIONES[0], level=1)
    _texto(doc, "resumen")

    _encargo(doc)

    doc.add_heading(SECCIONES[2], level=1)
    _ficha(
        doc,
        [
            ("Marca", "{{ ctx.marca }}"),
            ("Sector", "{{ ctx.sector }}"),
            ("Motores analizados", "{{ ctx.motores }}"),
            ("Fecha del estudio", "{{ ctx.fecha }}"),
            ("Origen de los datos", "{{ ctx.origen }}"),
            ("Datos sintéticos", "{{ ctx.sinteticos }}"),
            ("Estudio", "{{ ctx.study_id }}"),
        ],
    )
    doc.add_heading("Limitaciones", level=2)
    _vinetas(doc, "ctx.limitaciones")

    _diagnostico(doc)
    _competencia(doc)
    _hallazgos(doc)
    _exactitud(doc)

    doc.add_heading(SECCIONES[7], level=1)
    _texto(doc, "estrategia.texto")
    doc.add_heading("Objetivos", level=2)
    _tabla(
        doc,
        ["KPI", "Valor actual", "Valor objetivo", "Plazo"],
        "estrategia.objetivos",
        ["kpi", "actual", "objetivo", "plazo"],
    )
    doc.add_heading("Prioridades", level=2)
    _vinetas(doc, "estrategia.prioridades")

    _plan(doc)
    _escenarios(doc)
    _seguimiento(doc)

    doc.add_heading(SECCIONES[11], level=1)
    _tabla(doc, ["Riesgo", "Mitigación"], "riesgos", ["riesgo", "mitigacion"])

    doc.add_heading(SECCIONES[12], level=1)
    doc.add_heading("Hitos", level=2)
    _tabla(doc, ["Semana", "Hito"], "hoja.hitos", ["semana", "descripcion"])
    doc.add_heading("Responsabilidades", level=2)
    _tabla(doc, ["Quién", "Qué"], "hoja.responsabilidades", ["quien", "que"])
    _ficha(
        doc,
        [("Presupuesto", "{{ hoja.presupuesto }}"), ("Seguimiento", "{{ hoja.seguimiento }}")],
    )

    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(path))
    return path
